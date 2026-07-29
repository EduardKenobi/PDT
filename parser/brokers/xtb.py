import openpyxl
from . import xtb_old, xtb_new

def detect_xtb_format(filepath):
    """
    Detects if the file is in the new XTB Excel format or the old format.
    Returns 'new' or 'old'.
    """
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        sheetnames = wb.sheetnames
        # close read-only workbook to avoid file handle leaks
        wb.close()
        if 'Closed Positions' in sheetnames or 'Cash Operations' in sheetnames:
            return 'new'
    except Exception as e:
        print(f"Warning: Failed to detect Excel format for {filepath}: {e}")
    return 'old'

def split_accounts_by_format(accounts):
    """
    Splits accounts into new-format and old-format lists.
    """
    new_accounts = []
    old_accounts = []
    for account in accounts:
        if detect_xtb_format(account) == 'new':
            new_accounts.append(account)
        else:
            old_accounts.append(account)
    return new_accounts, old_accounts

def get_transactions_from_excel_files(accounts):
    """
    Dispatcher that detects the file formats of accounts, runs the appropriate
    parsers (new or old), and merges the parsed transactions.
    """
    new_accounts, old_accounts = split_accounts_by_format(accounts)
    
    # Initialize empty merged structures
    merged_transactions = {'companies': {}}
    
    # 1. Parse and merge old-format accounts
    if old_accounts:
        print(f"Running OLD XTB parser on {len(old_accounts)} account(s)...")
        old_data = xtb_old.get_transactions_from_excel_files(old_accounts)
        _merge_transactions(merged_transactions, old_data)
        
    # 2. Parse and merge new-format accounts
    if new_accounts:
        print(f"Running NEW XTB parser on {len(new_accounts)} account(s)...")
        new_data = xtb_new.get_transactions_from_excel_files(new_accounts)
        _merge_transactions(merged_transactions, new_data)
        
    # Sort all transactions by date for each company
    for ticker in merged_transactions['companies']:
        merged_transactions['companies'][ticker]['transactions'].sort(
            key=lambda x: x.get('close_date') or x.get('open_date')
        )
        
    return merged_transactions

def get_cash_operations_from_excel_files(accounts):
    """
    Dispatcher that detects the file formats of accounts, runs the appropriate
    parsers (new or old), and merges the parsed cash operations and dividends.
    """
    new_accounts, old_accounts = split_accounts_by_format(accounts)
    
    merged_dividends = {'companies': {}}
    merged_cash_ops = {'other_operations': []}
    
    # 1. Parse and merge old-format accounts
    if old_accounts:
        print(f"Running OLD XTB cash operations parser on {len(old_accounts)} account(s)...")
        old_divs, old_ops = xtb_old.get_cash_operations_from_excel_files(old_accounts)
        _merge_dividends(merged_dividends, old_divs)
        merged_cash_ops['other_operations'].extend(old_ops.get('other_operations', []))
        
    # 2. Parse and merge new-format accounts
    if new_accounts:
        print(f"Running NEW XTB cash operations parser on {len(new_accounts)} account(s)...")
        new_divs, new_ops = xtb_new.get_cash_operations_from_excel_files(new_accounts)
        _merge_dividends(merged_dividends, new_divs)
        merged_cash_ops['other_operations'].extend(new_ops.get('other_operations', []))
        
    # Sort dividends
    for ticker in merged_dividends['companies']:
        merged_dividends['companies'][ticker]['dividends'].sort(key=lambda x: x.get('date'))
        
    # Sort cash operations
    merged_cash_ops['other_operations'].sort(key=lambda x: (x.get('date') is None, x.get('date')))
    
    return merged_dividends, merged_cash_ops

def _merge_transactions(base_data, new_data):
    for ticker, data in new_data.get('companies', {}).items():
        if ticker not in base_data['companies']:
            base_data['companies'][ticker] = {'transactions': []}
        base_data['companies'][ticker]['transactions'].extend(data.get('transactions', []))

def _merge_dividends(base_data, new_data):
    for ticker, data in new_data.get('companies', {}).items():
        if ticker not in base_data['companies']:
            base_data['companies'][ticker] = {'dividends': []}
        base_data['companies'][ticker]['dividends'].extend(data.get('dividends', []))